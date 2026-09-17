"""Lectura y validación del xlsx del inventario (`ReporteInventario.xlsx`).

Puro: no toca la base de datos ni sabe de HTTP. Lo usan dos llamantes que no deben
divergir en una sola coma:

  · `tools/import_inventory_snapshot.py`         — el importador de terminal.
  · `olo.services.inventory_snapshot_import`     — el mismo importador, por API.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

if TYPE_CHECKING:
    from pathlib import Path

# ── Columnas del archivo ─────────────────────────────────────────────────────
#
# Se nombran por su ENCABEZADO y no por su posición: el reporte tiene 77 columnas y
# el orden ha cambiado entre versiones del WMS. Buscar por nombre falla ruidosamente
# —«falta la columna X»— mientras que fiarse del índice lee el campo equivocado en
# silencio, que es peor.
COL_UBICACION = "Ubicación"
COL_ARTICULO = "Artículo"
COL_DESCRIPCION = "Descripción"
COL_CANTIDAD = "Cantidad Unidades"
COL_CANTIDAD_TXT = "Cantidad Almacenaje"
COL_PALLET = "Pallet"
COL_COMPANIA = "Nombre Compañía"
COL_LOTE = "Lote"
COL_CADUCIDAD = "Fecha Caducidad"
COL_SITUACION_UBIC = "Situación Ubicación"
COL_ESTADO_UBIC = "Estado Ubicación"

OBLIGATORIAS = (COL_UBICACION,)

# Lo que se guarda en `raw`. NO son las 77 columnas: `raw` existe para poder auditar
# una línea sin volver al Excel, no para duplicar el reporte en la base. 41.055 filas
# por 77 campos serian ~90 MB de JSONB que nadie consulta.
EN_RAW = (
    COL_ARTICULO, COL_DESCRIPCION, COL_CANTIDAD, COL_CANTIDAD_TXT, COL_PALLET,
    COL_COMPANIA, COL_LOTE, COL_CADUCIDAD, COL_SITUACION_UBIC, COL_ESTADO_UBIC,
    "Tipo Pallet", "Referencia ERP", "Código Ean", "Zona Almacenaje",
    "Tipo Ubicación", "Familia", "SubFamilia", "Peso", "Situación",
    "Fecha Ubicación", "IdAlmacenamiento", "Contenedor",
)


class ColumnaFaltanteError(Exception):
    """El xlsx no trae alguna de las columnas obligatorias."""


def _texto(v: Any) -> str | None:
    """Normaliza a texto o `None`. Un `''` y un `None` son lo mismo aquí."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _cantidad(v: Any) -> float | None:
    """La cantidad, o `None` si no la hay.

    `None` y `0` son distintos: `0` es «hay una línea y su cantidad es cero», que el
    WMS produce de verdad; `None` es «el reporte no lo dice». Convertir uno en otro
    haría que un hueco sin dato pareciera vacío.
    """
    if v is None or v == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return n if n >= 0 else None


def _uom(txt: Any) -> str | None:
    """La unidad de medida, de «54 UD».

    El reporte trae la cantidad dos veces: numérica en `Cantidad Unidades` y con su
    unidad pegada en `Cantidad Almacenaje`. Se toma el número de la primera y la
    unidad de la segunda, en lugar de parsear la segunda entera: si el formato
    cambiara, se perdería la unidad pero no la cantidad.
    """
    s = _texto(txt)
    if not s:
        return None
    m = re.search(r"[A-Za-z]+\s*$", s)
    return m.group(0).strip().upper()[:12] if m else None


def _fecha(v: Any) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for f in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s[:19], f).date()
        except ValueError:
            continue
    return None


def _json_seguro(v: Any) -> Any:
    """Valores que `json.dumps` no sabe serializar: fechas, sobre todo."""
    if isinstance(v, datetime | date):
        return v.isoformat()
    return v


def leer_inventario(
    origen: Path | Any,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], datetime | None]:
    """Lee el archivo. Devuelve (líneas, rechazos, fecha del snapshot).

    `origen` es una ruta o cualquier objeto tipo-archivo binario que openpyxl
    acepte (un `BytesIO` con el cuerpo subido por HTTP, en el caso de la API).

    La fecha del snapshot se toma de la `Fecha Ubicación` MÁS RECIENTE del archivo, y
    no de `now()`. Es la diferencia entre «esta foto es del almacén el martes» y «este
    archivo se subió el jueves»: con `now()`, dos importaciones del mismo reporte
    tendrían fechas distintas y no habría forma de ordenar las fotos por antigüedad
    real.
    """
    wb = load_workbook(origen, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas_iter = ws.iter_rows(values_only=True)

    cab = [(_texto(c) or "") for c in next(filas_iter)]
    idx = {nombre: i for i, nombre in enumerate(cab)}
    faltan = [c for c in OBLIGATORIAS if c not in idx]
    if faltan:
        wb.close()
        raise ColumnaFaltanteError(
            f"El archivo no tiene la columna obligatoria {faltan!r}. "
            f"Columnas encontradas: {cab[:8]}…"
        )

    def val(fila: tuple[Any, ...], nombre: str) -> Any:
        i = idx.get(nombre)
        return None if i is None or i >= len(fila) else fila[i]

    lineas: list[dict[str, Any]] = []
    rechazos: list[dict[str, Any]] = []
    tomada: datetime | None = None

    for n, fila in enumerate(filas_iter, start=2):
        if fila is None or all(c is None or c == "" for c in fila):
            continue
        codigo = _texto(val(fila, COL_UBICACION))
        if not codigo:
            rechazos.append({"fila": n, "motivo": "sin codigo de ubicacion"})
            continue

        f_ubic = val(fila, "Fecha Ubicación")
        if isinstance(f_ubic, datetime) and (tomada is None or f_ubic > tomada):
            tomada = f_ubic

        lineas.append(
            {
                "codigo": codigo.upper(),
                "sku": _texto(val(fila, COL_ARTICULO)),
                "descripcion": _texto(val(fila, COL_DESCRIPCION)),
                "qty": _cantidad(val(fila, COL_CANTIDAD)),
                "uom": _uom(val(fila, COL_CANTIDAD_TXT)),
                "pallet": _texto(val(fila, COL_PALLET)),
                "compania": _texto(val(fila, COL_COMPANIA)),
                "lote": _texto(val(fila, COL_LOTE)),
                "caduca": _fecha(val(fila, COL_CADUCIDAD)),
                "raw": {
                    k: _json_seguro(val(fila, k))
                    for k in EN_RAW
                    if val(fila, k) not in (None, "")
                },
            }
        )

    wb.close()
    return lineas, rechazos, tomada
