"""Lectura y validación del xlsx del catálogo espacial (`ReporteUbicaciones.xlsx`).

Puro: no toca la base de datos ni sabe de HTTP. Lo usan dos llamantes que no deben
divergir en una sola coma:

  · `tools/import_spatial_catalog.py`      — el importador de terminal.
  · `olo.services.spatial_catalog_import`  — el mismo importador, por API.

Antes de esto las dos copias vivían en el script y una tercera habría sido la
API: dos reglas de validación que un día dejan de coincidir sin que nadie lo note
hasta que el mismo archivo se acepta por un lado y se rechaza por el otro.
"""

from __future__ import annotations

import json
import re
import unicodedata
from typing import TYPE_CHECKING, Any

from openpyxl import load_workbook

if TYPE_CHECKING:
    from pathlib import Path

# ── Encabezados esperados, por nombre Y posición ────────────────────────────
# Un archivo con columnas reordenadas se rechaza entero: durante el análisis un
# índice desplazado un puesto produjo resultados plausibles y falsos.
ENCABEZADOS = [
    "Id Almacenamiento", "Ubicación", "Preámbulo", "Referencia", "Columna",
    "Nivel", "Posición", "Tipo Ubicación", "Estado", "Situación",
    "Eje X", "Eje Y", "Eje Z", "Peso Máximo", "Zona Almacenaje",
    "Zona Picking", "Zona Trabajo Recurso", "Zona Trabajo Preparación",
    "Zona Cola Preparación", "Id Ubicación",
]

# Techo de plausibilidad de capacidad, en kg. Debe coincidir con
# `core.capacity_ceiling('weight_kg')`; el importador lo COMPRUEBA contra la
# base antes de escribir, porque dos umbrales copiados divergen en cuanto uno
# se toca.
#
# No es una lista de centinelas: el catálogo real usa seis grafías distintas de
# «sin límite» (1e5, 1e6, 9999999, 1e7, 99999999, 1e8) y enumerarlas fue el
# defecto que la migración 0058 corrigió. Ver su cabecera para los datos medidos.
TECHO_PESO_KG = 50000

# Estado del WMS → estado del ESPACIO. `OCUP` no entra: la ocupación es del
# snapshot de `wms`, no del estante (SPA-11 y SPA-12).
ESTADO_WMS = {"DISP": "available", "BLOQ": "blocked", "OCUP": "available"}

PATRON_CODIGO = re.compile(r"^[A-Z0-9][A-Z0-9._]*-C[0-9]{3}-N[0-9]{2}-[0-9]$")

_ACENTOS = str.maketrans("ÁÀÄÂÃÅÉÈËÊÍÌÏÎÓÒÖÔÕÚÙÜÛÑÇ ", "AAAAAAEEEEIIIIOOOOOUUUUNC_")


class CabecerasInesperadasError(Exception):
    """El xlsx no trae los encabezados esperados, en el orden esperado."""


class RechazoError(Exception):
    def __init__(self, motivo: str, campo: str = "") -> None:
        super().__init__(motivo)
        self.motivo = motivo
        self.campo = campo


def normalizar(bruto: str | None) -> str:
    """Misma regla que `core.normalize_spatial_code()`, y tiene que coincidir.

    El espacio pasa a `_`, NUNCA a `-`: el guion separa segmentos, y `PHA LO`
    convertido a `PHA-LO` daría cinco segmentos.
    """
    if bruto is None:
        return ""
    limpio = unicodedata.normalize("NFC", str(bruto).strip().upper()).translate(_ACENTOS)
    return re.sub(r"[^A-Z0-9._-]", "_", limpio)


def texto(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def entero(v: Any) -> int | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return int(float(str(v).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return None


def capacidad(v: Any) -> tuple[int | None, int | None]:
    """Devuelve (capacidad utilizable, valor crudo si se descartó).

    Un valor por encima del techo no es una capacidad, es «sin límite» escrito
    con un número. Se anula, pero se DEVUELVE también el crudo: el valor no
    dice nada de la ubicación y sí dice algo del WMS de origen.
    """
    n = entero(v)
    if n is None or n == 0:
        return None, None
    return (None, n) if n >= TECHO_PESO_KG else (n, None)


def leer_catalogo(origen: Path | Any) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Devuelve (filas válidas, rechazos). No tumba el lote por una fila mala.

    `origen` es una ruta o cualquier objeto tipo-archivo binario que openpyxl
    acepte (un `BytesIO` con el cuerpo subido por HTTP, en el caso de la API).
    """
    wb = load_workbook(origen, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)

    cabecera = [texto(c) for c in next(it)]
    if cabecera[: len(ENCABEZADOS)] != ENCABEZADOS:
        wb.close()
        faltan = [h for h in ENCABEZADOS if h not in (cabecera or [])]
        raise CabecerasInesperadasError(
            f"Encabezados inesperados. Se esperaban {len(ENCABEZADOS)} en este orden. "
            f"faltan o cambiaron de sitio: {faltan or 'ninguno, pero el orden difiere'}. "
            f"recibido: {cabecera[:6]}…"
        )
    idx = {c: i for i, c in enumerate(cabecera)}

    filas: list[dict[str, Any]] = []
    rechazos: list[dict[str, Any]] = []

    for n, cruda in enumerate(it, start=2):
        try:
            ref_bruta = texto(cruda[idx["Referencia"]])
            ubic_bruta = texto(cruda[idx["Ubicación"]])
            if not ref_bruta:
                raise RechazoError("Referencia vacía", "Referencia")
            if not ubic_bruta:
                raise RechazoError("Ubicación vacía", "Ubicación")

            col = entero(cruda[idx["Columna"]])
            niv = entero(cruda[idx["Nivel"]])
            pos = entero(cruda[idx["Posición"]])
            if col is None or niv is None or pos is None:
                raise RechazoError("Columna, Nivel o Posición ausente", "Columna/Nivel/Posición")
            if not (1 <= niv <= 99):
                raise RechazoError(f"Nivel {niv} fuera de 1..99", "Nivel")
            if not (1 <= pos <= 9):
                raise RechazoError(f"Posición {pos} fuera de 1..9", "Posición")
            if col < 1:
                raise RechazoError(f"Columna {col} inválida", "Columna")

            ref_norm = normalizar(ref_bruta)
            code = f"{ref_norm}-C{col:03d}-N{niv:02d}-{pos}"

            # La clasificación no puede mentir: si no cumple el patrón, es `opaque`
            # y el parser estructurado NO se le aplica.
            forma = "structured" if PATRON_CODIGO.match(code) else "opaque"

            peso_max, peso_crudo = capacidad(cruda[idx["Peso Máximo"]])

            filas.append(
                {
                    "fila": n,
                    "ref_bruta": ref_bruta,
                    "ref_norm": ref_norm,
                    "storage_id": texto(cruda[idx["Id Almacenamiento"]]),
                    "preambulo": texto(cruda[idx["Preámbulo"]]),
                    "col": col, "niv": niv, "pos": pos,
                    "code": code,
                    "external_code": ubic_bruta,
                    "external_location_id": texto(cruda[idx["Id Ubicación"]]),
                    "tipo_wms": texto(cruda[idx["Tipo Ubicación"]]),
                    "estado_wms": texto(cruda[idx["Estado"]]),
                    "situacion": texto(cruda[idx["Situación"]]),
                    "lx": entero(cruda[idx["Eje X"]]),
                    "ly": entero(cruda[idx["Eje Y"]]),
                    "lz": entero(cruda[idx["Eje Z"]]),
                    "peso_max": peso_max,
                    "peso_max_crudo": peso_crudo,
                    "zona": texto(cruda[idx["Zona Almacenaje"]]),
                    "forma": forma,
                    "crudo": {
                        k: (str(cruda[i]) if cruda[i] is not None else None)
                        for k, i in idx.items() if k
                    },
                }
            )
        except RechazoError as r:
            rechazos.append(
                {"fila": n, "motivo": r.motivo, "campo": r.campo,
                 "crudo": json.dumps([str(c) if c is not None else None for c in cruda[:8]])}
            )

    wb.close()
    return filas, rechazos


def agregar(filas: list[dict[str, Any]]) -> tuple[dict[str, dict[str, Any]], set[tuple[str, int]]]:
    """Agregados: qué nodos (racks, cuerpos) hace falta crear a partir de las filas.

    Misma lógica que la sección `── Agregados ──` de `import_spatial_catalog.py`,
    factorizada porque la usan tanto el importador de terminal como el de la API.
    """
    refs: dict[str, dict[str, Any]] = {}
    for f in filas:
        r = refs.setdefault(
            f["ref_norm"],
            {"externo": f["ref_bruta"], "storage_id": f["storage_id"],
             "preambulo": f["preambulo"], "cols": set(), "tipos": set(), "zonas": set()},
        )
        r["cols"].add(f["col"])
        if f["tipo_wms"]:
            r["tipos"].add(f["tipo_wms"])
        if f["zona"]:
            r["zonas"].add(f["zona"])

    bays = {(f["ref_norm"], f["col"]) for f in filas}
    return refs, bays
